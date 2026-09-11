# -*- coding: utf-8 -*-
"""
motor/excel_newcontasis.py — Genera el Excel de importación del Registro de Compras de NewContaSis
("FORMATO REGISTRO DE COMPRAS - SISTEMA EXPERTO CONTABLE 25.00") y el Excel de revisión.

Mapa de columnas (según la plantilla oficial y los comentarios de sus celdas):
 A fecha emisión           B fecha vencimiento/pago     C tipo doc (01/03/07/08)   D serie (6)
 E año DUA                 F número (13)                G tipo doc proveedor (6/1)  H número doc (11)
 I razón social (60)       J base gravada→gravadas      K IGV                       L/M gravadas mixtas
 N/O gravadas→no gravadas  P no gravadas                Q ISC                       R otros tributos
 S importe total           T N° doc no domiciliado      U N° constancia detracción  V fecha depósito
 W tipo de cambio (10,4)   X fecha doc original         Y tipo                      Z serie
 AA número                 AB moneda S/D                AC equivalente USD          AD fecha vencimiento
 AE CON/CRE                AF cuenta base imponible     AG cuenta otros tributos    AH cuenta total
 AI/AJ centros de costo    AK régimen (1 det/2 perc/3 ret) AL % régimen             AM importe régimen
 AN serie doc régimen      AO número doc régimen        AP fecha doc régimen        AQ código presupuesto
 AR % IGV                  AS glosa (60)                AT condición percepción 1/2 AU importe base régimen
 AV clasificación 1-5      AW ICBPER                    AX cuenta ICBPER
El archivo de importación lleva SOLO filas de datos (la plantilla indica eliminar las filas 1-13).
"""
from __future__ import annotations

import io
from datetime import datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .clasificador import Propuesta

CABECERAS = ['FECHA EMISION', 'FECHA VENC/PAGO', 'TIPO DOC', 'SERIE', 'AÑO DUA', 'NUMERO', 'TIPO DOC PROV', 'NUM DOC PROV', 'RAZON SOCIAL',
             'BASE IMP GRAV', 'IGV', 'BASE IMP MIXTA', 'IGV MIXTA', 'BASE IMP NO GRAV', 'IGV NO GRAV', 'VALOR NO GRAVADO', 'ISC', 'OTROS TRIBUTOS', 'IMPORTE TOTAL',
             'N° DOC NO DOMIC', 'N° CONSTANCIA DETRAC', 'FECHA DEPOSITO', 'TIPO CAMBIO', 'FECHA DOC ORIG', 'TIPO DOC ORIG', 'SERIE DOC ORIG', 'NUM DOC ORIG',
             'MONEDA', 'EQUIV USD', 'FECHA VENCIMIENTO', 'CONDICION', 'CTA BASE IMPONIBLE', 'CTA OTROS TRIBUTOS', 'CTA TOTAL', 'CENTRO COSTOS', 'CENTRO COSTOS 2',
             'REGIMEN ESP', '% REGIMEN', 'IMPORTE REGIMEN', 'SERIE DOC REG', 'NUM DOC REG', 'FECHA DOC REG', 'COD PRESUPUESTO', '% IGV', 'GLOSA',
             'COND PERCEPCION', 'IMPORTE BASE REGIMEN', 'CLASIF BIENES/SERV', 'ICBPER', 'CTA ICBPER']
assert len(CABECERAS) == 50  # A..AX


def _f(iso: str):
    """'2026-08-17' → datetime (Excel fecha) o None."""
    try:
        return datetime.strptime(iso[:10], '%Y-%m-%d')
    except Exception:
        return None


def _n(v) -> float:
    return float(Decimal(str(v or 0)).quantize(Decimal('0.01')))


def _serie_numero(sn: str) -> tuple[str, str]:
    if '-' in sn:
        s, n = sn.split('-', 1)
        return s.strip()[:6], n.strip()[:13]
    return sn[:6], ''


def fila_newcontasis(p: Propuesta, cuenta_haber: str, cuenta_icbper: str = '', cuenta_otros_tributos: str = '') -> list:
    c = p.c
    serie, numero = _serie_numero(c.serie_numero)
    f_emi = _f(c.fecha_emision)
    f_ven = _f(c.fecha_vencimiento) or f_emi
    es_usd = c.moneda_codigo == 'USD'
    forma = (c.forma_pago or '').upper()
    condicion = 'CRE' if ('CRED' in forma or c.cuotas) else 'CON'
    tipo_prov = '6' if len(c.ruc_emisor) == 11 else ('1' if len(c.ruc_emisor) == 8 else '0')

    regimen = ''
    pct = imp = base_reg = ''
    if c.tiene_detraccion:
        regimen, pct, imp, base_reg = 1, _n(c.detraccion_porcentaje), _n(c.detraccion_monto), _n(c.importe_total)
    elif c.tiene_percepcion:
        regimen, pct, imp, base_reg = 2, _n(c.percepcion_porcentaje), _n(c.percepcion_monto), _n(c.importe_total)
    elif c.tiene_retencion:
        regimen, imp, base_reg = 3, _n(c.retencion_monto), _n(c.importe_total)

    no_grav = _n(c.exoneradas + c.inafectas)
    otros_trib = _n(c.otros_tributos + c.otros_cargos)
    ref_tipo = ref_serie = ref_num = ''
    if c.tipo_codigo in ('07', '08') and c.doc_referencia:
        ref_tipo = '01' if c.doc_referencia[:1].upper() == 'F' else ('03' if c.doc_referencia[:1].upper() == 'B' else '01')
        ref_serie, ref_num = _serie_numero(c.doc_referencia)

    return [
        f_emi, f_ven, c.tipo_codigo, serie, '', numero, tipo_prov, c.ruc_emisor[:11], c.nombre_emisor[:60],   # A-I
        _n(c.gravadas), _n(c.igv), 0.0, 0.0, 0.0, 0.0, no_grav, _n(c.isc), otros_trib, _n(c.importe_total),  # J-S
        '', '', None,                                                                                          # T-V  (constancia de detracción: no viene en el XML)
        (p.tc if es_usd else 1.0) or 1.0,                                                                      # W
        None, ref_tipo, ref_serie, ref_num,                                                                     # X-AA
        'D' if es_usd else 'S', _n(c.importe_total) if es_usd else '', f_ven, condicion,                        # AB-AE
        p.cuenta, (cuenta_otros_tributos or p.cuenta) if otros_trib else '', cuenta_haber, '', '',              # AF-AJ
        regimen, pct, imp, '', '', None, '',                                                                    # AK-AQ
        18.0 if c.igv else 0.0, p.glosa[:60],                                                                   # AR-AS
        '1' if c.tiene_percepcion else '', base_reg,                                                            # AT-AU
        p.av, _n(c.icbper) if c.icbper else '', (cuenta_icbper if c.icbper else ''),                            # AV-AX
    ]


FORMATOS = {1: 'dd/mm/yyyy', 2: 'dd/mm/yyyy', 22: 'dd/mm/yyyy', 23: '#,##0.0000', 24: 'dd/mm/yyyy', 30: 'dd/mm/yyyy', 42: 'dd/mm/yyyy'}
NUMERICAS = set(range(10, 20)) | {29, 38, 39, 44, 47, 49}
TEXTO = {3, 4, 5, 6, 7, 8, 9, 20, 21, 25, 26, 27, 28, 31, 32, 33, 34, 35, 36, 40, 41, 43, 45, 46, 48, 50}


def _escribir_filas(ws, filas, inicio=1):
    for i, fila in enumerate(filas):
        for j, v in enumerate(fila, start=1):
            cell = ws.cell(row=inicio + i, column=j, value=v)
            if j in FORMATOS and v is not None:
                cell.number_format = FORMATOS[j]
            elif j in NUMERICAS and v != '':
                cell.number_format = '#,##0.00'
            elif j in TEXTO:
                cell.number_format = '@'


def generar_excel_importacion(props: list[Propuesta], cuenta_haber: str, cuenta_icbper: str = '', con_cabecera: bool = False) -> bytes:
    """Excel listo para importar: solo filas de datos (o con cabecera de referencia si con_cabecera)."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'COMPRAS'
    filas = [fila_newcontasis(p, cuenta_haber, cuenta_icbper) for p in props if p.estado in ('ok', 'revisar') and p.cuenta]
    inicio = 1
    if con_cabecera:
        ws.append(CABECERAS)
        inicio = 2
    _escribir_filas(ws, filas, inicio)
    for j in range(1, 51):
        ws.column_dimensions[get_column_letter(j)].width = 14
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generar_excel_revision(props: list[Propuesta], cuenta_haber: str, empresa: str, periodo: str) -> bytes:
    """Excel amigable: hoja PROPUESTA (todas las filas con semáforo), REVISAR, EXCLUIDOS, y NEWCONTASIS (con cabecera)."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'PROPUESTA'
    verde, amarillo, rojo, gris = 'C6EFCE', 'FFEB9C', 'FFC7CE', 'E7E6E6'
    cab = ['Semáforo', 'Confianza', 'Estado', 'Fecha', 'Tipo', 'Serie-Número', 'RUC', 'Proveedor', 'Moneda', 'Base', 'IGV', 'Total', 'TC',
           'Cuenta', 'Descripción cuenta', 'Cuenta haber', 'Clasif. NCS', 'Fuente', 'Detracción', 'Glosa', 'Alertas', 'Explicación', 'Archivo']
    ws.append([f'{empresa} — Propuesta contable de compras — {periodo}'])
    ws['A1'].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(cab)
    for cell in ws[3]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1F3B5C')
    for p in props:
        c = p.c
        fila = [p.semaforo, p.confianza if p.estado in ('ok', 'revisar') else '', p.estado.upper() + (f' — {p.motivo}' if p.motivo else ''),
                c.fecha_emision, c.tipo_codigo, c.serie_numero, c.ruc_emisor, c.nombre_emisor, c.moneda_codigo, _n(c.gravadas) or _n(c.valor_venta), _n(c.igv), _n(c.importe_total),
                p.tc if p.tc else '', p.cuenta, p.cuenta_desc, cuenta_haber if p.cuenta else '', p.av if p.cuenta else '', p.fuente,
                f'{c.detraccion_porcentaje}% S/ {c.detraccion_monto}' if c.tiene_detraccion else '', p.glosa, ' | '.join(p.alertas), p.explicacion, c.archivo]
        ws.append(fila)
        color = {'🟢': verde, '🟡': amarillo, '🔴': rojo, '⚪': gris}[p.semaforo]
        for cell in ws[ws.max_row][:3]:
            cell.fill = PatternFill('solid', fgColor=color)
    anchos = [9, 10, 22, 11, 6, 16, 13, 36, 8, 11, 10, 11, 8, 11, 40, 11, 9, 10, 16, 40, 60, 60, 30]
    for j, w in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = 'A4'

    ws2 = wb.create_sheet('REVISAR')
    ws2.append(['Semáforo', 'Serie-Número', 'Proveedor', 'Total', 'Cuenta propuesta', 'Confianza', 'Motivo / alertas'])
    for cell in ws2[1]:
        cell.font = Font(bold=True)
    for p in props:
        if p.estado == 'revisar' or (p.estado == 'ok' and (p.alertas or p.confianza < 85)):
            ws2.append([p.semaforo, p.c.serie_numero, p.c.nombre_emisor, _n(p.c.importe_total), p.cuenta, p.confianza, ' | '.join(p.alertas) or p.explicacion])
    for j, w in enumerate([9, 16, 36, 11, 12, 10, 90], start=1):
        ws2.column_dimensions[get_column_letter(j)].width = w

    ws3 = wb.create_sheet('EXCLUIDOS')
    ws3.append(['Estado', 'Serie-Número', 'RUC', 'Proveedor', 'Total', 'Motivo', 'Archivo'])
    for cell in ws3[1]:
        cell.font = Font(bold=True)
    for p in props:
        if p.estado in ('excluido', 'duplicado', 'error'):
            ws3.append([p.estado, p.c.serie_numero, p.c.ruc_emisor, p.c.nombre_emisor, _n(p.c.importe_total), p.motivo, p.c.archivo])
    for j, w in enumerate([11, 16, 13, 36, 11, 60, 30], start=1):
        ws3.column_dimensions[get_column_letter(j)].width = w

    ws4 = wb.create_sheet('NEWCONTASIS (ref)')
    ws4.append(CABECERAS)
    for cell in ws4[1]:
        cell.font = Font(bold=True, size=8)
        cell.alignment = Alignment(wrap_text=True)
    _escribir_filas(ws4, [fila_newcontasis(p, cuenta_haber) for p in props if p.estado in ('ok', 'revisar') and p.cuenta], 2)
    for j in range(1, 51):
        ws4.column_dimensions[get_column_letter(j)].width = 13
    ws4.freeze_panes = 'A2'

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
