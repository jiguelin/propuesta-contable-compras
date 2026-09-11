# -*- coding: utf-8 -*-
"""Pruebas del motor (sin Streamlit, sin IA):  python -m pytest -q"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl  # noqa: E402

from motor.excel_newcontasis import CABECERAS  # noqa: E402
from motor.plan_cuentas import Cuenta, PlanCuentas  # noqa: E402
from motor.servicio import Servicio  # noqa: E402
from motor.tipo_cambio import HistorialTC, leer_pdf_sunat  # noqa: E402
from motor.xml_parser import parse_xml  # noqa: E402

DATOS = Path(__file__).parent / 'datos'


def plan_minimo() -> bytes:
    """Plan de cuentas genérico (cuenta | descripción) en Excel."""
    import io
    wb = openpyxl.Workbook()
    ws = wb.active
    for cod, desc in [('4212', 'EMITIDAS'), ('1011', 'CAJA MN'), ('6343094', 'MANTENIMIENTO PPE - ADM'), ('6343093', 'MANTENIMIENTO PPE - CDS'),
                      ('603202521', 'SUMINISTROS COMBUSTIBLES'), ('6560094', 'SUMINISTROS - ADM'), ('6391094', 'GASTOS BANCARIOS - ADM'), ('33611', 'EQUIPO DE COMPUTO')]:
        ws.append([cod, desc])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parser_detraccion_usd():
    c = parse_xml(DATOS / '20100000001-01-F001-000123.xml')
    assert not c.error and c.tipo_codigo == '01' and c.moneda_codigo == 'USD'
    assert c.tiene_detraccion and float(c.detraccion_porcentaje) == 12 and c.detraccion_cuenta == '00-000-123456'
    assert float(c.igv) == 180 and float(c.importe_total) == 1180 and c.forma_pago == 'Credito'


def test_tipo_cambio_pdf_y_fallback():
    tc = HistorialTC(leer_pdf_sunat((DATOS / 'tipocambio_2026-05.pdf').read_bytes()))
    assert len(tc) == 31
    t, exacto = tc.buscar('2026-05-17')
    assert exacto and t.venta == 3.435
    t, exacto = tc.buscar('2026-06-02')      # sin publicación → último publicado (31/05)
    assert not exacto and t.fecha.day == 31


def test_plan_busqueda_y_hojas():
    p = PlanCuentas([Cuenta('63', 'GASTOS'), Cuenta('634', 'MANT'), Cuenta('6343094', 'MANTENIMIENTO ADM')])
    p.reindex()
    assert not p.get('63').es_hoja and p.get('6343094').es_hoja
    assert [c.codigo for c in p.con_prefijo('634')] == ['6343094']
    assert p.buscar('manten')[0].codigo in ('634', '6343094')


def test_flujo_completo_y_aprendizaje(tmp_path):
    s = Servicio(ruta_db=tmp_path / 'p.db')
    s.registrar_empresa('20611889683', 'ALQUMIN EIRL', plan_minimo())
    s.cargar_tc_pdf((DATOS / 'tipocambio_2026-05.pdf').read_bytes())
    archivos = [(f.name, f.read_bytes()) for f in DATOS.glob('*.xml')]
    lote = s.procesar('20611889683', archivos, cuenta_haber='4212', usar_ia=False, excluir_detraccion_sin_constancia=False)
    por_ruc = {p.c.ruc_emisor: p for p in lote.propuestas}
    banco, serv = por_ruc['20100047218'], por_ruc['20100000001']
    assert banco.estado == 'excluido' and 'Banco' in banco.motivo
    assert serv.estado in ('ok', 'revisar') and serv.cuenta.startswith('6343') and serv.tc == 3.435
    assert any('detracción' in a for a in serv.alertas)

    # Excel de importación: 2 filas (servicio USD + laptop), 50 columnas, cuenta del haber en AH, régimen 1 en AK, moneda D
    wb = openpyxl.load_workbook(filename=_bytes_io(lote.excel_importacion()))
    ws = wb.active
    assert ws.max_row == 2 and ws.max_column == len(CABECERAS) == 50
    fila = next([c.value for c in r] for r in ws.iter_rows() if r[7].value == '20100000001')
    assert fila[2] == '01' and fila[3] == 'F001' and fila[5] == '000123' and fila[27] == 'D' and fila[22] == 3.435
    assert fila[31] == serv.cuenta and fila[33] == '4212' and fila[36] == 1 and fila[30] == 'CRE'

    # corrección manual → memoria → segundo lote sale con alta confianza
    s.aplicar_correccion(lote, serv.id, '6343093')
    lote2 = s.procesar('20611889683', archivos, cuenta_haber='1011', usar_ia=False, excluir_detraccion_sin_constancia=False)
    serv2 = {p.c.ruc_emisor: p for p in lote2.propuestas}['20100000001']
    assert serv2.cuenta == '6343093' and serv2.fuente == 'memoria' and serv2.confianza >= 92
    assert 'RESUMEN' in lote2.reporte_txt() and lote2.zip_todo()


def _bytes_io(b: bytes):
    import io
    return io.BytesIO(b)


def test_periodo_constancias_activo_y_usd(tmp_path):
    from datetime import date
    from motor.tipo_cambio import TC
    from motor.excel_newcontasis import fila_newcontasis
    s = Servicio(ruta_db=tmp_path / 'p.db')
    s.registrar_empresa('20611889683', 'ALQUMIN EIRL', plan_minimo())
    s.db.guardar_tc([TC(date(2026, 5, d), 3.50, 3.52) for d in range(1, 32)])
    archivos = [(f.name, f.read_bytes()) for f in DATOS.glob('*.xml')]
    txt = ('Numero Constancia|RUC Proveedor|Fecha Pago|Monto Deposito|Serie de Comprobante|Numero de Comprobante|\n'
           '312739411|20100000001|20/05/2026|420.00|F001|00000123|\n999|20999999999|20/05/2026|1.00|F001|1|\n').encode()
    lote = s.procesar('20611889683', archivos, cuenta_haber='4212', usar_ia=False, periodo='2026-05', constancias=[('c.txt', txt)])
    por_ruc = {p.c.ruc_emisor: p for p in lote.propuestas}
    serv = por_ruc['20100000001']
    assert serv.det_constancia == '312739411' and serv.det_fecha == date(2026, 5, 20)
    assert lote.n_constancias == 1 and [c.numero for c in lote.constancias_sin_factura] == ['999']
    fila = fila_newcontasis(serv, '4212')
    assert fila[20] == '312739411' and fila[22] == 3.52 and fila[27] == 'D'
    assert fila[9] == 3520.0 and fila[18] == round(1180 * 3.52, 2) and fila[28] == 1180.0   # USD → soles en J..S, USD en AC
    lap = por_ruc['20100000002']
    assert lap.posible_activo and any('ACTIVO FIJO' in a for a in lap.alertas)
    assert fila_newcontasis(lap, '4212')[22] == 3.52                                          # TC también para soles
    # sin constancia → la factura con detracción queda fuera del Excel (regla tributaria)
    lote3 = s.procesar('20611889683', archivos, cuenta_haber='4212', usar_ia=False, periodo='2026-05')
    serv3 = {p.c.ruc_emisor: p for p in lote3.propuestas}['20100000001']
    assert serv3.estado == 'excluido' and 'constancia' in serv3.motivo
    import openpyxl, io
    ws = openpyxl.load_workbook(io.BytesIO(lote3.excel_importacion())).active
    assert ws.max_row == 1 and ws['H1'].value == '20100000002' and ws['W1'].value == 3.52
